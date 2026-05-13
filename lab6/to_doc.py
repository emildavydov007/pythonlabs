from docx import Document
def save_to_doc(area, heat):
    doc = Document()
    doc.add_heading('Отчёт', 0)
    doc.add_paragraph(f'Площадь: {area} м²')
    doc.add_paragraph(f'Мощность: {heat} Вт')
    doc.save("result.docx")